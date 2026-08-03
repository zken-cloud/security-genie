#!/usr/bin/env python3
"""Build Security Genie slide decks (PPTX) from docs/.

Decks:
  slides/security-genie-workshop.pptx  <- docs/facilitator-guide.md
  slides/gcp-setup.pptx                <- docs/gcp-setup.md

Import to Google Slides: upload the .pptx to Google Drive, then
"Open with > Google Slides" (or File > Import slides from an existing deck).

Regenerate after editing this file or the docs:
  python3 -m venv /tmp/slides-venv && /tmp/slides-venv/bin/pip install python-pptx
  /tmp/slides-venv/bin/python slides/build_slides.py
"""
import os

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

OUT_DIR = os.path.dirname(os.path.abspath(__file__))

BLUE = RGBColor(0x1A, 0x73, 0xE8)
DARK = RGBColor(0x20, 0x24, 0x28)
GRAY = RGBColor(0x5F, 0x63, 0x68)
CODE = RGBColor(0x0B, 0x3D, 0x91)

BODY_FONT = "Roboto"       # native to Google Slides
MONO_FONT = "Roboto Mono"  # native to Google Slides


def new_deck():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    return prs


def _box(slide, x, y, w, h):
    box = slide.shapes.add_textbox(x, y, w, h)
    box.text_frame.word_wrap = True
    return box.text_frame


def _run(p, text, size, bold=False, color=DARK, mono=False):
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color
    r.font.name = MONO_FONT if mono else BODY_FONT


def title_slide(prs, title, subtitle):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    tf = _box(slide, Inches(0.9), Inches(2.7), Inches(11.5), Inches(2.2))
    _run(tf.paragraphs[0], title, 42, bold=True, color=BLUE)
    _run(tf.add_paragraph(), subtitle, 18, color=GRAY)


def content_slide(prs, title, items):
    """items: ('h', text) sub-header | ('b', text[, level]) bullet |
              ('t', text) plain | ('c', text) code line"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _run(_box(slide, Inches(0.6), Inches(0.35), Inches(12.1), Inches(0.9))
         .paragraphs[0], title, 28, bold=True, color=BLUE)
    tf = _box(slide, Inches(0.65), Inches(1.35), Inches(12.0), Inches(5.8))
    first = True
    for item in items:
        kind = item[0]
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        if kind == "h":
            p.space_before = Pt(12)
            _run(p, item[1], 19, bold=True)
        elif kind == "b":
            level = item[2] if len(item) > 2 else 0
            p.level = level
            p.space_after = Pt(5)
            _run(p, ("•  " if level == 0 else "–  ") + item[1], 16)
        elif kind == "t":
            p.space_after = Pt(5)
            _run(p, item[1], 16)
        elif kind == "c":
            p.space_after = Pt(1)
            _run(p, item[1], 13, mono=True, color=CODE)


def table_slide(prs, title, headers, rows, widths=None, note=None, font=12):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _run(_box(slide, Inches(0.6), Inches(0.35), Inches(12.1), Inches(0.9))
         .paragraphs[0], title, 28, bold=True, color=BLUE)
    shape = slide.shapes.add_table(len(rows) + 1, len(headers),
                                   Inches(0.6), Inches(1.4),
                                   Inches(12.1), Inches(0.4) * (len(rows) + 1))
    table = shape.table
    if widths:
        for i, w in enumerate(widths):
            table.columns[i].width = Inches(w)
    for c, text in enumerate(headers):
        cell = table.cell(0, c)
        _run(cell.text_frame.paragraphs[0], text, font + 1, bold=True, color=BLUE)
    for r, row in enumerate(rows, start=1):
        for c, text in enumerate(row):
            cell = table.cell(r, c)
            cell.text_frame.word_wrap = True
            _run(cell.text_frame.paragraphs[0], text, font)
    if note:
        _run(_box(slide, Inches(0.65), Inches(6.7), Inches(12.0), Inches(0.6))
             .paragraphs[0], note, 12, color=GRAY)


# ---------------------------------------------------------------- deck 1
def build_workshop():
    prs = new_deck()

    title_slide(prs, "Security Genie",
                "A security-specialist harness + skill pack for coding agents — facilitator deck")

    content_slide(prs, "The program: four sessions", [
        ("b", "1. Build + connect (60 min) — harness tour, GCP MCP setup, multi-account credentials, guard hooks"),
        ("b", "2. Upgrade part 1 (45 min) — well-architected review, least-privilege Terraform review + generator"),
        ("b", "3. Upgrade part 2 (45 min) — security jargon translator, security objection handling"),
        ("b", "4. Genie at work (60 min) — flawed architecture end-to-end: review, fix, document, hand over"),
        ("h", "Works on"),
        ("b", "Claude Code  •  Google Antigravity (agy)  •  OpenAI Codex  •  Gemini CLI"),
    ])

    content_slide(prs, "What the harness is", [
        ("b", "One AGENTS.md persona: senior Google Cloud platform security specialist"),
        ("b", "Nine portable skills in the cross-agent SKILL.md format"),
        ("b", "Live GCP access via the official gcloud MCP server (@google-cloud/gcloud-mcp)"),
        ("b", "Multi-account credential tooling (named gcloud configurations + scripts/gcp-use.sh)"),
        ("b", "Pre-exec guard hooks that block dangerous commands until the user consents"),
        ("b", "A deliberately flawed example architecture for practice (examples/flawed-agent-app)"),
    ])

    content_slide(prs, "Design: one canonical instruction file", [
        ("b", "AGENTS.md holds the persona and operating rules exactly once"),
        ("b", "CLAUDE.md and GEMINI.md are two-line adapters that import it"),
        ("b", "Codex and Antigravity read AGENTS.md natively"),
        ("h", "Why"),
        ("b", "Four agents, one source of truth — edit the rules in one place, every agent follows"),
    ])

    content_slide(prs, "Design: portable SKILL.md + trigger descriptions", [
        ("b", "SKILL.md is the cross-agent standard: YAML frontmatter (name, description) + markdown body"),
        ("b", "Canonical copies live in skills/ and are linked into .claude/skills, .agents/skills, .gemini/skills"),
        ("b", "The description field is the trigger: it names the task AND the phrases users actually say"),
        ("h", "Why"),
        ("b", "Skills survive an agent switch — no vendor lock-in for workshop material"),
        ("b", "When a skill fails to fire in a session, the description is the first thing to fix"),
    ])

    content_slide(prs, "Design: evidence before opinion", [
        ("b", "Every review gathers evidence first: read-only gcloud MCP for live environments, files for IaC"),
        ("b", "Every finding must cite file:line or a resource ID — no evidence, no finding"),
        ("b", "(VERIFY) marks uncertain role/constraint names instead of inventing plausible ones"),
        ("h", "Why"),
        ("b", "A review that hallucinates findings destroys customer trust faster than no review"),
        ("b", "Models confidently invent IAM role names and org-policy constraint IDs — make uncertainty explicit"),
    ])

    content_slide(prs, "Design: hooks are the backstop, not the boundary", [
        ("b", "AGENTS.md instructs read-only; the guard hook enforces it mid-turn"),
        ("b", "Blocks: rm -rf class, package/binary installs, pipe-to-shell, terraform apply/destroy,"),
        ("b", "sudo, destructive git (force push, reset --hard, clean -f), kubectl delete, dd/mkfs", 1),
        ("b", "Isolated installs stay allowed (venv / pipx) — the sanctioned path"),
        ("b", "Fails open on bad input — a hook bug must never wedge a session"),
        ("h", "Why"),
        ("b", "Instructions drift over long contexts and sub-agents; a deterministic pre-exec check does not"),
        ("b", "It is a consent-enforcing speed bump, not a sandbox — say this aloud in session 1"),
    ])

    table_slide(prs, "The nine skills", ["Skill", "Job"], [
        ("security-faq", "Fast, defensible answers to recurring GCP security questions"),
        ("threat-model", "STRIDE-per-element over GCP trust boundaries"),
        ("security-review", "Read-only audit of Terraform, designs, or live projects"),
        ("well-architected-review", "Architecture Framework assessment, maturity 1-5, roadmap"),
        ("terraform-least-privilege-review", "Pattern-matched IAM audit (P1-P10) with minimal-diff fixes"),
        ("terraform-secure-generator", "Secure-by-default Terraform + security decision log"),
        ("sast-scan", "Semgrep / Checkov / Trivy / tfsec wrapper + triage"),
        ("security-jargon-translator", "Security terms translated per audience"),
        ("security-objection-handling", "Prep sheets for customer security objections"),
    ], widths=[3.6, 8.5], font=13)

    content_slide(prs, "Why: security-faq and threat-model", [
        ("h", "security-faq"),
        ("b", "CEs answer the same ~15 questions on every engagement — a position cheat sheet"),
        ("b", "plus verification commands keeps answers consistent and checkable", 1),
        ("h", "threat-model"),
        ("b", "GCP incidents are almost always boundary crossings: internet edge, project,"),
        ("b", "VPC SC perimeter, CI/CD supply chain, identities — the cheat sheet focuses there", 1),
        ("b", "Output: threat table with existing controls, gaps, GCP-native mitigations, mermaid DFD"),
    ])

    content_slide(prs, "Why: security-review and well-architected-review", [
        ("h", "security-review"),
        ("b", "The bread-and-butter deliverable: domain checklist of real misconfigs + severity rubric"),
        ("b", "The rubric forces prioritization over laundry lists", 1),
        ("h", "well-architected-review"),
        ("b", "The business-level opener: pillar scores (1-5) turn issues into a fundable roadmap"),
        ("b", "Deep security findings defer to security-review — the two never diverge", 1),
    ])

    content_slide(prs, "Why: terraform-least-privilege-review and terraform-secure-generator", [
        ("h", "terraform-least-privilege-review"),
        ("b", "IAM is the #1 source of GCP findings; the same ten violations recur everywhere"),
        ("b", "P1-P10 codify them as grep-able patterns with ready minimal-diff remediations", 1),
        ("h", "terraform-secure-generator"),
        ("b", "Flagging flaws is half the job — the generator produces the fix artifact"),
        ("b", "The security decision log doubles as handover documentation", 1),
    ])

    content_slide(prs, "Why: sast-scan", [
        ("b", "Automates the breadth pass: Semgrep for code, Checkov/Trivy/tfsec for Terraform"),
        ("b", "Triage turns scanner hits into verified findings: dedupe, reachability, severity, FP justification"),
        ("h", "The lesson built in"),
        ("b", "Scanners are not reviews — every scan ends with a mandatory 'Not covered' list"),
        ("b", "Proof point in the example: semgrep catches eval and shell=True but misses the planted SQLi"),
    ])

    content_slide(prs, "Why: security-jargon-translator and security-objection-handling", [
        ("h", "security-jargon-translator"),
        ("b", "Customer conversations fail on vocabulary, not facts"),
        ("b", "Per-audience framing: exec = money/risk, engineer = mechanism, auditor = control/evidence", 1),
        ("h", "security-objection-handling"),
        ("b", "Objections are predictable: residency, CLOUD Act, 'cloud is less secure', cost"),
        ("b", "Framework: acknowledge > clarify > respond > honest trade-offs > next step", 1),
        ("b", "The honest-limits section is what builds trust", 1),
    ])

    content_slide(prs, "Facilitating: what good looks like", [
        ("b", "Grade agent output against each skill's Output template — deviation means the skill wasn't loaded"),
        ("h", "Common pitfalls"),
        ("b", "Agent answers from memory instead of loading the skill > name the skill explicitly once"),
        ("b", "Agent skips evidence on live questions > check MCP connection and the active gcloud profile"),
        ("b", "Findings without file:line evidence > send them back; the skills forbid it"),
    ])

    table_slide(prs, "Deployment map", ["Agent", "Harness", "Skills", "MCP", "Hooks"], [
        ("Claude Code", "CLAUDE.md > AGENTS.md", ".claude/skills/", ".mcp.json", ".claude/settings.json (PreToolUse)"),
        ("Antigravity (agy)", "AGENTS.md", ".agents/skills/", ".gemini/settings.json", ".gemini/settings.json (BeforeTool)"),
        ("Codex", "AGENTS.md", ".agents/skills/", ".codex/config.toml", "~/.codex/hooks.json (experimental)"),
        ("Gemini CLI", "GEMINI.md > AGENTS.md", ".gemini/skills/", ".gemini/settings.json", ".gemini/settings.json (BeforeTool)"),
    ], widths=[2.1, 2.6, 2.1, 2.5, 2.8], font=11,
        note="Codex hooks need [features] hooks = true. agy hook location: VERIFY against the installed version.")

    content_slide(prs, "Installing and updating", [
        ("c", "./install.sh --agent all --target /path/to/repo   # customer / workshop repo"),
        ("c", "./install.sh --agent claude                       # one agent only"),
        ("c", "./install.sh --agent all --global                 # every repo, for your user"),
        ("c", "./install.sh --link --target /path/to/repo        # symlinks: edit once, live everywhere"),
        ("h", "Semantics"),
        ("b", "Project scope for workshops/customer repos; global for a CE's daily driver"),
        ("b", "Re-running replaces skill copies but never clobbers existing settings — merge hints instead"),
    ])

    content_slide(prs, "Guard hook: deploy and test", [
        ("c", "python3 hooks/dangerous_command_guard.py --selftest     # 43 cases"),
        ("c", "echo '{\"tool_input\":{\"command\":\"rm -rf /\"}}' | python3 .claude/hooks/dangerous_command_guard.py"),
        ("c", "# exit=2 (blocked)    terraform plan -> exit=0 (allowed)"),
        ("h", "End-to-end demo"),
        ("b", "Ask the agent to 'delete everything and reinstall' in a test repo — it gets blocked and asks for consent"),
        ("h", "Limitations to state"),
        ("b", "Pattern-matches shell commands only; cannot catch danger via allowed commands; Codex hooks experimental"),
    ])

    content_slide(prs, "Authoring a new skill", [
        ("b", "mkdir skills/<name> — directory name must equal the frontmatter name"),
        ("b", "description: 1-2 sentences naming the task AND the user phrases that should trigger it"),
        ("b", "Sections: When to use / Inputs (<=3 questions) / Workflow (evidence first) / content / Output format / Guardrails"),
        ("b", "90-150 lines; dense beats long; (VERIFY) on uncertain names; never invent rule/role IDs"),
        ("b", "Register in AGENTS.md and README.md skill tables, then re-run install.sh (or use --link)"),
    ])

    content_slide(prs, "Exercise: examples/flawed-agent-app", [
        ("b", "Scenario: 'customer-provided' Acme support agent — Cloud Run, Cloud SQL, HTTP LB, public bucket, bastion"),
        ("b", "Order: sast-scan > terraform-least-privilege-review > security-review > threat-model > fix with terraform-secure-generator"),
        ("b", "Score against EXPECTED-FINDINGS.md: 27 planted flaws (7 Critical / 9 High / 9 Medium / 2 Low)"),
        ("b", "Score substance over row-count matches — scanners and skills overlap by design"),
        ("h", "Capstone"),
        ("b", "Deploy the FIXED Terraform to a smoke-test project; never deploy the flawed version"),
    ])

    table_slide(prs, "Troubleshooting", ["Symptom", "Likely cause / fix"], [
        ("Skill never triggers", "Not installed for this agent/scope, or weak description — fix and re-install"),
        ("MCP tools missing", "Wrong config file, untrusted project (Codex), or Node/npx unavailable — docs/gcp-setup.md"),
        ("Answers about the wrong project", "ADC is global — re-run scripts/gcp-use.sh <profile> --adc and re-confirm identity"),
        ("Hook never fires", "Settings not merged, wrong event name (PreToolUse vs BeforeTool), or Codex flag off"),
        ("Hook blocks a legit command", "Ruleset too broad — edit hooks/dangerous_command_guard.py, re-run --selftest"),
    ], widths=[3.4, 8.7], font=12)

    prs.save(os.path.join(OUT_DIR, "security-genie-workshop.pptx"))


# ---------------------------------------------------------------- deck 2
def build_gcp_setup():
    prs = new_deck()

    title_slide(prs, "GCP setup for Security Genie",
                "gcloud MCP server + multi-account credential management")

    content_slide(prs, "How it works", [
        ("b", "The genie uses the official Google Cloud MCP server: @google-cloud/gcloud-mcp"),
        ("b", "It wraps the local gcloud CLI — the agent gets whatever identity gcloud is configured with"),
        ("h", "Consequence"),
        ("b", "Credential hygiene IS access control: wrong profile or stale ADC = agent reads the wrong customer"),
    ])

    content_slide(prs, "Prerequisites", [
        ("c", "gcloud --version          # Google Cloud CLI"),
        ("c", "node --version            # Node 18+ (for npx)"),
        ("c", "gcloud auth login"),
        ("c", "gcloud auth application-default login"),
        ("c", "npx -y @google-cloud/gcloud-mcp --help     # smoke test"),
    ])

    table_slide(prs, "MCP config per agent", ["Agent", "Project config", "User-level config"], [
        ("Claude Code", ".mcp.json", "claude mcp add gcloud -- npx -y @google-cloud/gcloud-mcp"),
        ("Gemini CLI", ".gemini/settings.json", "npx @google-cloud/gcloud-mcp init --agent=gemini-cli"),
        ("Antigravity (agy)", ".gemini/settings.json (when supported)", "~/.gemini/antigravity/mcp_config.json"),
        ("Codex", ".codex/config.toml (trusted projects)", "~/.codex/config.toml"),
    ], widths=[2.3, 4.1, 5.7], font=11,
        note="Verify after setup: claude mcp list / gemini mcp list / codex mcp list.")

    content_slide(prs, "The server block (same in every format)", [
        ("c", "command: npx"),
        ("c", "args:    [\"-y\", \"@google-cloud/gcloud-mcp\"]"),
        ("c", "env:     CLOUDSDK_ACTIVE_CONFIG_NAME = <profile>   # optional profile pin"),
        ("h", "Notes"),
        ("b", "Codex reads project config only for trusted projects — otherwise merge into ~/.codex/config.toml"),
        ("b", "Other official Google Cloud MCP servers (observability, Cloud Run, ...) can be added the same way"),
    ])

    content_slide(prs, "Multi-account model", [
        ("b", "Named gcloud configurations are the profile mechanism"),
        ("b", "One per customer/account/environment: acme-prod, acme-dev, globex-sandbox", 1),
        ("b", "Each pins an account and a default project"),
        ("h", "The trap"),
        ("b", "ADC (Application Default Credentials) is a single global file — it does NOT follow the active configuration"),
        ("b", "This is the #1 cause of 'agent looked at the wrong customer'", 1),
    ])

    content_slide(prs, "Create and switch profiles", [
        ("c", "gcloud config configurations create acme-prod"),
        ("c", "gcloud config set account you@example.com --configuration=acme-prod"),
        ("c", "gcloud config set project  acme-prod-123   --configuration=acme-prod"),
        ("c", "gcloud auth login --configuration=acme-prod"),
        ("h", "Day to day"),
        ("c", "scripts/gcp-use.sh list                  # profiles + current identity"),
        ("c", "source scripts/gcp-use.sh acme-prod      # activate + export CLOUDSDK_ACTIVE_CONFIG_NAME"),
    ])

    content_slide(prs, "ADC: two workable patterns", [
        ("h", "1. User ADC, refreshed on every switch"),
        ("c", "source scripts/gcp-use.sh acme-prod --adc"),
        ("b", "Simple; easy to forget", 1),
        ("h", "2. Impersonated service account per customer (recommended)"),
        ("c", "gcloud auth application-default login --impersonate-service-account=sec-review@acme-prod-123.iam.gserviceaccount.com"),
        ("b", "Needs roles/iam.serviceAccountTokenCreator on that SA; consistent until reset", 1),
        ("h", "Either way"),
        ("b", "State the account and project in every review report"),
    ])

    content_slide(prs, "Verify before trusting", [
        ("c", "gcloud config get-value account"),
        ("c", "gcloud config get-value project"),
        ("c", "gcloud auth list --filter=status:ACTIVE"),
        ("c", "gcloud config configurations list"),
        ("h", "MCP side"),
        ("c", "claude mcp list   |   gemini mcp list   |   codex mcp list"),
    ])

    content_slide(prs, "Troubleshooting", [
        ("b", "npx ... gcloud-mcp fails > Node/npm missing, or proxy blocks the registry"),
        ("b", "MCP connected but wrong project > profile/ADC mismatch — run scripts/gcp-use.sh list"),
        ("b", "Codex ignores .codex/config.toml > project not trusted; merge into ~/.codex/config.toml"),
        ("b", "Permission-denied from the agent is usually correct behavior:"),
        ("b", "grant the review identity viewer roles (roles/viewer, roles/iam.securityReviewer, roles/logging.viewer)", 1),
    ])

    prs.save(os.path.join(OUT_DIR, "gcp-setup.pptx"))


if __name__ == "__main__":
    build_workshop()
    build_gcp_setup()
    print("wrote security-genie-workshop.pptx and gcp-setup.pptx to", OUT_DIR)
